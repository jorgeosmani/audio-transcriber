"""
exportar.py
-----------
Da formato a las transcripciones y las guarda en TXT, Markdown, Word (.docx)
o en un archivo por audio (.txt o .srt).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from motor import Resultado, Segmento

_LINEA = "═" * 64

_ORIGENES = {
    "nombre": "tomada del nombre del archivo",
    "modificacion": "fecha de modificación (el nombre no tenía fecha)",
    "sin_fecha": "sin fecha",
}


# ------------------------------------------------------------------ utilidades
def tiempo(segundos: float, srt: bool = False) -> str:
    """65.3 -> '01:05'  |  3725 -> '1:02:05'  |  srt=True -> '00:01:05,300'"""
    segundos = max(0.0, segundos)
    h, resto = divmod(int(segundos), 3600)
    m, s = divmod(resto, 60)
    if srt:
        ms = int(round((segundos - int(segundos)) * 1000))
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def texto_fecha(r: Resultado) -> str:
    if r.fecha is None:
        return "sin fecha"
    formato = "%d/%m/%Y %H:%M" if r.tiene_hora else "%d/%m/%Y"
    return f"{r.fecha.strftime(formato)}  ({_ORIGENES.get(r.origen_fecha, r.origen_fecha)})"


def encabezado(numero: int, r: Resultado) -> str:
    detalles = []
    if r.duracion:
        detalles.append(f"Duración: {tiempo(r.duracion)}")
    if r.idioma:
        detalles.append(f"Idioma: {r.idioma}")
    lineas = [
        _LINEA,
        f"[{numero}] {r.ruta.name}",
        f"Fecha: {texto_fecha(r)}",
    ]
    if detalles:
        lineas.append("   ".join(detalles))
    lineas.append(_LINEA)
    return "\n".join(lineas) + "\n"


def linea_segmento(seg: Segmento, marcas: bool) -> str:
    """Con marcas: una línea por frase. Sin marcas: texto corrido."""
    if marcas:
        return f"[{tiempo(seg.inicio)} → {tiempo(seg.fin)}] {seg.texto}\n"
    return seg.texto + " "


def cuerpo(r: Resultado, marcas: bool) -> str:
    if r.error:
        return f"⚠ No se pudo transcribir este archivo: {r.error}\n"
    if not r.segmentos:
        return "(No se detectó voz en este audio)\n"
    if marcas:
        return "".join(linea_segmento(s, True) for s in r.segmentos)
    return r.texto + "\n"


def texto_combinado(resultados: list[Resultado], marcas: bool = True) -> str:
    partes = []
    for i, r in enumerate(resultados, 1):
        partes.append(encabezado(i, r) + cuerpo(r, marcas))
    return "\n".join(partes)


# ------------------------------------------------------------------- guardado
def guardar_txt(resultados: list[Resultado], destino: str | Path, marcas: bool = True) -> Path:
    destino = Path(destino)
    cabecera = (
        f"Transcripción generada el {datetime.now():%d/%m/%Y %H:%M}\n"
        f"Archivos: {len(resultados)} (ordenados por fecha)\n\n"
    )
    destino.write_text(cabecera + texto_combinado(resultados, marcas), encoding="utf-8")
    return destino


def guardar_markdown(resultados: list[Resultado], destino: str | Path, marcas: bool = True) -> Path:
    destino = Path(destino)
    lineas = [
        "# Transcripción de audios",
        "",
        f"Generada el {datetime.now():%d/%m/%Y %H:%M} · {len(resultados)} archivos ordenados por fecha.",
        "",
    ]
    for i, r in enumerate(resultados, 1):
        lineas += [f"## {i}. {r.ruta.name}", "", f"**Fecha:** {texto_fecha(r)}  "]
        if r.duracion:
            lineas.append(f"**Duración:** {tiempo(r.duracion)}  ")
        if r.idioma:
            lineas.append(f"**Idioma:** {r.idioma}  ")
        lineas.append("")
        if r.error:
            lineas.append(f"> ⚠ No se pudo transcribir: {r.error}")
        elif not r.segmentos:
            lineas.append("_(No se detectó voz en este audio)_")
        elif marcas:
            lineas += [f"- `{tiempo(s.inicio)}` {s.texto}" for s in r.segmentos]
        else:
            lineas.append(r.texto)
        lineas.append("")
    destino.write_text("\n".join(lineas), encoding="utf-8")
    return destino


def guardar_docx(resultados: list[Resultado], destino: str | Path, marcas: bool = True) -> Path:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError as e:
        raise ImportError(
            "Para exportar a Word instala python-docx:  pip install python-docx"
        ) from e

    destino = Path(destino)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Transcripción de audios", level=0)
    doc.add_paragraph(
        f"Generada el {datetime.now():%d/%m/%Y %H:%M}. "
        f"{len(resultados)} archivos ordenados por fecha."
    )

    gris = RGBColor(0x66, 0x66, 0x66)
    for i, r in enumerate(resultados, 1):
        doc.add_heading(f"{i}. {r.ruta.name}", level=1)
        meta = [f"Fecha: {texto_fecha(r)}"]
        if r.duracion:
            meta.append(f"Duración: {tiempo(r.duracion)}")
        if r.idioma:
            meta.append(f"Idioma: {r.idioma}")
        p = doc.add_paragraph()
        run = p.add_run("   |   ".join(meta))
        run.italic = True
        run.font.color.rgb = gris

        if r.error:
            doc.add_paragraph(f"⚠ No se pudo transcribir: {r.error}")
        elif not r.segmentos:
            doc.add_paragraph("(No se detectó voz en este audio)")
        elif marcas:
            for s in r.segmentos:
                p = doc.add_paragraph()
                marca = p.add_run(f"[{tiempo(s.inicio)}]  ")
                marca.bold = True
                marca.font.color.rgb = gris
                p.add_run(s.texto)
        else:
            doc.add_paragraph(r.texto)

    doc.save(destino)
    return destino


def _srt(r: Resultado) -> str:
    bloques = []
    for i, s in enumerate(r.segmentos, 1):
        bloques.append(f"{i}\n{tiempo(s.inicio, True)} --> {tiempo(s.fin, True)}\n{s.texto}\n")
    return "\n".join(bloques)


def guardar_individuales(
    resultados: list[Resultado], carpeta: str | Path, formato: str = "txt", marcas: bool = True
) -> int:
    """
    Guarda un archivo por audio en la carpeta indicada. Los nombres llevan un
    prefijo numérico (001_, 002_, …) para conservar el orden cronológico.
    Devuelve cuántos archivos se escribieron.
    """
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    escritos = 0
    for i, r in enumerate(resultados, 1):
        if r.error:
            continue
        nombre = f"{i:03d}_{r.ruta.stem}.{formato}"
        if formato == "srt":
            contenido = _srt(r)
        else:
            contenido = encabezado(i, r) + cuerpo(r, marcas)
        (carpeta / nombre).write_text(contenido, encoding="utf-8")
        escritos += 1
    return escritos
